from pathlib import Path

from docx import Document

from third_edition_work_import import load_builder


builder = load_builder()


def test_builder_renders_heading_code_table_and_body(tmp_path: Path) -> None:
    chapter = tmp_path / "chapter.md"
    chapter.write_text(
        "# 第一章 示例\n\n正文。\n\n## 1.1 小节\n\n```bash\necho ok\n```\n\n"
        "| 项目 | 值 |\n| --- | --- |\n| A | B |\n",
        encoding="utf-8",
    )
    output = tmp_path / "book.docx"
    builder.build([chapter], output)
    doc = Document(output)
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    assert "第一章 示例" in text
    assert "echo ok" in text
    assert len(doc.tables) == 1
    assert doc.tables[0].cell(1, 1).text == "B"


def test_output_is_not_source_markdown(tmp_path: Path) -> None:
    chapter = tmp_path / "chapter.md"
    chapter.write_text("# 第一章 示例\n", encoding="utf-8")
    output = tmp_path / "book.docx"
    builder.build([chapter], output)
    assert output.read_bytes()[:2] == b"PK"
    assert output.read_bytes() != chapter.read_bytes()
