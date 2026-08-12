from __future__ import annotations

import argparse
import csv
import hashlib
import re
import zipfile
from collections import Counter
from pathlib import Path

from docx import Document


HEADING = re.compile(r"^Heading ([12])$")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def extract(source: Path, work_root: Path) -> None:
    research = work_root / "research"
    media_root = work_root / "assets" / "second-edition-media"
    research.mkdir(parents=True, exist_ok=True)
    media_root.mkdir(parents=True, exist_ok=True)

    doc = Document(source)
    headings: list[tuple[int, int, str]] = []
    for index, paragraph in enumerate(doc.paragraphs):
        match = HEADING.match(paragraph.style.name)
        if match and paragraph.text.strip():
            headings.append((index, int(match.group(1)), paragraph.text.strip()))

    ranges: list[dict[str, object]] = []
    for pos, (start, level, title) in enumerate(headings):
        end = headings[pos + 1][0] if pos + 1 < len(headings) else len(doc.paragraphs)
        texts = [p.text.strip() for p in doc.paragraphs[start:end] if p.text.strip()]
        disposition = "retain-and-update"
        reason = "保留知识主线并按第三版事实和实训环境更新。"
        if "终端软件" in title:
            disposition, reason = "delete", "按用户要求删除独立终端工具教学，仅保留一页选型说明。"
        elif "安装脚本" in title:
            disposition, reason = "delete", "第二篇统一改为逐条手工命令与配置，不再把脚本解读作为安装方法。"
        elif "Windows" in title or "CentOS" in " ".join(texts[:20]):
            disposition, reason = "rewrite", "替换过时来宾系统与产品界面，主用openEuler SP3/CirrOS。"
        ranges.append(
            {
                "level": level,
                "title": title,
                "start_paragraph": start,
                "end_paragraph": end - 1,
                "nonempty_paragraphs": len(texts),
                "characters": sum(len(text) for text in texts),
                "disposition": disposition,
                "reason": reason,
            }
        )

    map_path = research / "second-edition-content-map.csv"
    with map_path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(ranges[0]))
        writer.writeheader()
        writer.writerows(ranges)

    style_counts = Counter(p.style.name for p in doc.paragraphs if p.text.strip())
    summary = research / "second-edition-structure.md"
    summary.write_text(
        "\n".join(
            [
                "# 第二版原稿结构盘点",
                "",
                f"- 源文件：`{source.name}`",
                f"- SHA-256：`{sha256(source)}`",
                f"- 段落：{len(doc.paragraphs)}",
                f"- 表格：{len(doc.tables)}",
                f"- 节：{len(doc.sections)}",
                f"- 一级标题：{style_counts['Heading 1']}",
                f"- 二级标题：{style_counts['Heading 2']}",
                "- Word 实测页数：377（2026-08-12，Microsoft Word 重新分页）。",
                "- 处理原则：删除终端工具独立章节和重复脚本解读；保留可复用理论与操作图，所有活动安装步骤改写为openEuler SP3 + OpenStack Antelope纯手工流程。",
                "",
                "详细范围见 `second-edition-content-map.csv`。",
            ]
        )
        + "\n",
        encoding="utf-8",
    )

    with zipfile.ZipFile(source) as archive:
        media = sorted(name for name in archive.namelist() if name.startswith("word/media/") and not name.endswith("/"))
        for name in media:
            target = media_root / Path(name).name
            target.write_bytes(archive.read(name))

    manifest = research / "second-edition-media.csv"
    with manifest.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(["source_name", "local_path", "bytes", "sha256"])
        for path in sorted(media_root.iterdir(), key=lambda item: item.name.lower()):
            writer.writerow([path.name, path.relative_to(work_root).as_posix(), path.stat().st_size, sha256(path)])


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("work_root", type=Path)
    args = parser.parse_args()
    extract(args.source.resolve(), args.work_root.resolve())


if __name__ == "__main__":
    main()
