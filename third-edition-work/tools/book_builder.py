from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK, WD_LINE_SPACING, WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


HEADING_RE = re.compile(r"^(#{1,4})\s+(.+)$")
IMAGE_RE = re.compile(r"^!\[([^]]*)]\(([^)]+)\)$")
TABLE_SEPARATOR_RE = re.compile(r"^\s*\|?\s*:?-{3,}")


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def ensure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.first_line_indent = Pt(21)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    for name, size, color in (
        ("Heading 1", 22, "17365D"),
        ("Heading 2", 16, "1F4E79"),
        ("Heading 3", 13, "2F5597"),
        ("Heading 4", 11, "385723"),
    ):
        style = styles[name]
        style.font.name = "黑体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
    if "Code Block" not in styles:
        style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Consolas"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
        style.font.size = Pt(8.5)
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.line_spacing = 1.0
    if "Figure Caption" not in styles:
        style = styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "宋体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style.font.size = Pt(9)
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    if "Part Title" not in styles:
        style = styles.add_style("Part Title", WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "黑体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(24)
        style.font.bold = True
        style.paragraph_format.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER


def add_page_number(section) -> None:
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def configure_sections(doc: Document) -> None:
    for section in doc.sections:
        section.top_margin = Cm(1.27)
        section.bottom_margin = Cm(1.27)
        section.left_margin = Cm(1.27)
        section.right_margin = Cm(1.27)
        section.header_distance = Cm(0.6)
        section.footer_distance = Cm(0.6)
        add_page_number(section)


def add_code_paragraph(doc: Document, line: str) -> None:
    paragraph = doc.add_paragraph(style="Code Block")
    paragraph.add_run(line or " ")
    properties = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F2F2F2")
    properties.append(shading)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if len(rows) < 2:
        return
    header, body = rows[0], rows[2:]
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    for index, text in enumerate(header):
        cell = table.rows[0].cells[index]
        cell.text = text.strip()
        set_cell_shading(cell, "D9EAF7")
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in body:
        cells = table.add_row().cells
        for index in range(len(header)):
            cells[index].text = row[index].strip() if index < len(row) else ""


def parse_table_row(line: str) -> list[str]:
    value = line.strip()
    if value.startswith("|"):
        value = value[1:]
    if value.endswith("|"):
        value = value[:-1]
    return [part.strip() for part in value.split("|")]


def append_markdown(doc: Document, path: Path) -> None:
    lines = path.read_text(encoding="utf-8").splitlines()
    index = 0
    in_code = False
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if stripped.startswith("```"):
            in_code = not in_code
            index += 1
            continue
        if in_code:
            add_code_paragraph(doc, line)
            index += 1
            continue
        heading = HEADING_RE.match(line)
        if heading:
            level = len(heading.group(1))
            title = heading.group(2).strip()
            if title.startswith("第") and "篇" in title:
                doc.add_paragraph(title, style="Part Title")
            else:
                doc.add_heading(title, level=min(level, 4))
            index += 1
            continue
        if stripped.startswith("|") and index + 1 < len(lines) and TABLE_SEPARATOR_RE.match(lines[index + 1]):
            rows = [parse_table_row(line), parse_table_row(lines[index + 1])]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(parse_table_row(lines[index]))
                index += 1
            add_table(doc, rows)
            continue
        image = IMAGE_RE.match(stripped)
        if image:
            source = (path.parent / image.group(2)).resolve()
            if source.exists():
                paragraph = doc.add_paragraph()
                paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
                paragraph.add_run().add_picture(str(source), width=Cm(14.4))
                if image.group(1):
                    doc.add_paragraph(image.group(1), style="Figure Caption")
            else:
                doc.add_paragraph(f"[图片待插入：{image.group(1)}]", style="Figure Caption")
            index += 1
            continue
        if stripped.startswith(("- ", "* ")):
            paragraph = doc.add_paragraph(stripped[2:].strip(), style="List Bullet")
            index += 1
            continue
        if re.match(r"^\d+\.\s+", stripped):
            paragraph = doc.add_paragraph(re.sub(r"^\d+\.\s+", "", stripped), style="List Number")
            index += 1
            continue
        if stripped.startswith(">"):
            paragraph = doc.add_paragraph(stripped.lstrip("> "), style="Intense Quote")
            index += 1
            continue
        if stripped:
            doc.add_paragraph(stripped)
        index += 1


def new_book(title: str) -> Document:
    doc = Document()
    ensure_styles(doc)
    configure_sections(doc)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    paragraph.paragraph_format.space_before = Pt(180)
    run = paragraph.add_run(title)
    run.bold = True
    run.font.name = "黑体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    run.font.size = Pt(28)
    subtitle = doc.add_paragraph("第三版初稿")
    subtitle.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    subtitle.runs[0].font.size = Pt(18)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading("编写说明", level=1)
    doc.add_paragraph("本稿资料事实冻结于2026年7月31日。概览采用OpenStack 2026.1 Gazpacho；实训采用openEuler 24.03 LTS SP3上的OpenStack 2023.1 Antelope。")
    doc.add_paragraph("第二篇所有部署和运维命令均由学习者逐条手工输入；第十三章是明确区分的智能体自动化实训。")
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading("目录", level=1)
    doc.add_paragraph("请在Microsoft Word中右键更新目录域。")
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    return doc


def build(chapter_paths: list[Path], output: Path) -> None:
    doc = new_book("云计算基础架构平台构建与应用")
    for position, path in enumerate(chapter_paths):
        append_markdown(doc, path)
        if position + 1 < len(chapter_paths):
            doc.add_section(WD_SECTION.NEW_PAGE)
    normalize_figure_placeholders(doc)
    configure_sections(doc)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def normalize_figure_placeholders(doc: Document) -> None:
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("[绘图建议：") or text.startswith("[截图待采集："):
            paragraph.style = doc.styles["Figure Caption"]


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("output", type=Path)
    parser.add_argument("chapters", nargs="+", type=Path)
    args = parser.parse_args()
    build([path.resolve() for path in args.chapters], args.output.resolve())


if __name__ == "__main__":
    main()
